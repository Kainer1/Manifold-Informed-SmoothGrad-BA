"""Export radial Fourier statistics of saved AllSamples images to Excel."""

import argparse
from contextlib import ExitStack
from pathlib import Path
import sys

import h5py
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from Helper.paths import RESULTS_ROOT  # noqa: E402

RGB_MEAN = np.array([0.485, 0.456, 0.406], dtype=np.float64)[None, :, None, None]
RGB_STD = np.array([0.229, 0.224, 0.225], dtype=np.float64)[None, :, None, None]


def radial_bins(size):
    """Retain the original BA radial bins and their coefficient counts."""
    frequency = np.fft.fftfreq(size)
    radius = np.hypot(frequency[:, None], frequency[None, :])
    bins = np.floor(radius * size + 1e-10).astype(int)
    valid = (bins > 0) & (bins < size // 2)
    bin_ids = bins[valid]
    counts = np.bincount(bin_ids, minlength=size // 2)[1:]
    centers = (
        np.bincount(bin_ids, weights=radius[valid], minlength=size // 2)[1:] / counts
    )
    return valid, bin_ids, counts, centers


def radial_spectra(normalized_samples, valid, bin_ids, counts):
    """Return mean amplitude and total power per ring for each input sample."""
    if not np.isfinite(normalized_samples).all():
        raise ValueError("Samples contain non-finite values.")
    rgb = normalized_samples.astype(np.float64) * RGB_STD + RGB_MEAN
    gray = rgb.mean(axis=1)
    gray -= gray.mean(axis=(-2, -1), keepdims=True)
    fft = np.fft.fft2(gray, axes=(-2, -1), norm="forward")
    amplitude = np.abs(fft)
    power = np.abs(fft) ** 2
    amplitudes = np.empty((len(normalized_samples), len(counts)), dtype=np.float64)
    powers = np.empty_like(amplitudes)
    for position in range(len(normalized_samples)):
        amplitudes[position] = (
            np.bincount(
                bin_ids,
                weights=amplitude[position][valid],
                minlength=len(counts) + 1,
            )[1:]
            / counts
        )
        powers[position] = np.bincount(
            bin_ids, weights=power[position][valid], minlength=len(counts) + 1
        )[1:]
    if not np.isfinite(amplitudes).all() or not np.isfinite(powers).all():
        raise ValueError("Fourier spectra contain non-finite values.")
    return amplitudes, powers


def open_samples(stack, directory, sampler):
    if not sampler or Path(sampler).name != sampler or sampler in {".", ".."}:
        raise ValueError("Supply an exact sampler name without a path.")
    source = stack.enter_context(h5py.File(directory / f"samples_{sampler}.h5", "r"))
    for key, expected in (
        ("sampler_name", sampler),
        ("model_name", "VGG16"),
        ("dataset_name", "IMAGENET_256"),
    ):
        if source.attrs.get(key) != expected:
            raise ValueError(
                f"Invalid {key} in {source.filename}; expected {expected}."
            )
    if source.attrs.get("representation", "normalized model input") != (
        "normalized model input"
    ):
        raise ValueError(f"Expected normalized model inputs in {source.filename}.")
    return source


def check_group(source, index):
    key = str(index)
    if key not in source:
        raise ValueError(f"Missing image index {index} in {source.filename}.")
    group = source[key]
    if not bool(group.attrs.get("complete", False)):
        raise ValueError(f"Incomplete image index {index} in {source.filename}.")
    samples = group["samples"]
    if samples.ndim != 4 or samples.shape[1:] != (3, 256, 256) or not len(samples):
        raise ValueError(f"Invalid sample shape at image index {index}.")
    if samples.dtype != np.dtype("float32"):
        raise ValueError(f"Expected float32 samples at image index {index}.")
    if int(group.attrs.get("sample_count", -1)) != len(samples):
        raise ValueError(f"Sample count mismatch at image index {index}.")
    label = int(group.attrs["label"])
    if not 0 <= label < 1000:
        raise ValueError(f"Invalid label at image index {index}.")
    return samples, label


def analyze(args):
    if (args.index is None) != (args.sample_index is None):
        raise ValueError("Use --index and --sample-index together, or omit both.")
    if args.index is not None and (args.index < 0 or args.sample_index < 0):
        raise ValueError("Image indices and sample positions must be nonnegative.")
    if args.batch_size < 1:
        raise ValueError("Batch size must be positive.")
    quantiles = sorted(args.quantiles)
    if not quantiles or any(not 0 <= q <= 1 for q in quantiles):
        raise ValueError("Quantiles must lie between 0 and 1.")
    if len(set(quantiles)) != len(quantiles):
        raise ValueError("Quantiles must not contain duplicates.")
    base = args.results_dir / "VGG16" / "IMAGENET_256"
    directory = base / "Samples" / "allSamples"
    valid, bin_ids, counts, centers = radial_bins(256)
    amplitude_batches, power_batches, relative_batches, coverage = [], [], [], []
    with ExitStack() as stack:
        source = open_samples(stack, directory, args.sampler)
        reference = (
            open_samples(stack, directory, args.reference_sampler)
            if args.reference_sampler
            else None
        )
        if args.index is None:
            keys = [key for key in source if key != "_pending"]
            if any(not key.isdigit() or str(int(key)) != key for key in keys):
                raise ValueError(
                    "All image-index group names must be nonnegative integers."
                )
            indices = sorted(map(int, keys))
        else:
            indices = [args.index]
        if not indices:
            raise ValueError("No stored image indices found.")
        for index in indices:
            samples, label = check_group(source, index)
            first = 0 if args.sample_index is None else args.sample_index
            stop = len(samples) if args.sample_index is None else first + 1
            if first >= len(samples):
                raise ValueError(
                    f"Sample position {first} is missing at image {index}."
                )
            if reference is not None:
                original, original_label = check_group(reference, index)
                if len(original) != 1 or original_label != label:
                    raise ValueError(
                        f"Reference needs one sample and the same label at {index}."
                    )
                _, original_power = radial_spectra(original[:], valid, bin_ids, counts)
                if not np.all(original_power > 0):
                    raise ValueError(f"Zero reference power at image index {index}.")
            for start in range(first, stop, args.batch_size):
                batch = samples[start : min(start + args.batch_size, stop)]
                amplitude, power = radial_spectra(batch, valid, bin_ids, counts)
                amplitude_batches.append(amplitude)
                power_batches.append(power / counts)
                if reference is not None:
                    relative_batches.append(100 * (power / original_power - 1))
            coverage.append((index, label, stop - first, first, stop - 1))
            print(f"Processed image {index}: {stop - first} samples", flush=True)
        source_path = source.filename
        reference_path = reference.filename if reference is not None else ""

    metrics = {
        "Amplitude": np.concatenate(amplitude_batches),
        "Power": np.concatenate(power_batches),
    }
    if relative_batches:
        metrics["RelativePowerPercent"] = np.concatenate(relative_batches)
    sample_count = len(metrics["Power"])
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    headers = ["CyclesPerPixel", "CoefficientsInRing"]
    columns = [centers, counts]
    for metric, values in metrics.items():
        headers.append(f"{metric}Mean")
        columns.append(values.mean(axis=0))
        for quantile in quantiles:
            headers.append(f"{metric}P{quantile * 100:g}")
            columns.append(np.quantile(values, quantile, axis=0, method="linear"))
    summary.append(headers)
    for row in zip(*columns):
        summary.append([value.item() for value in row])
    sheet = workbook.create_sheet("Coverage")
    sheet.append(
        [
            "ImageIndex",
            "GroundTruthLabel",
            "SampleCount",
            "FirstSamplePosition",
            "LastSamplePosition",
        ]
    )
    for row in coverage:
        sheet.append(row)
    definitions = workbook.create_sheet("Definitions")
    definitions.append(["Setting", "Value"])
    for row in (
        ("Sampler", args.sampler),
        ("Source", source_path),
        ("ReferenceSampler", args.reference_sampler or ""),
        ("ReferenceSource", reference_path),
        ("ImageCount", len(coverage)),
        ("SampleCount", sample_count),
        (
            "Selection",
            "all stored indices and samples" if args.index is None else "single sample",
        ),
        ("Quantiles", ", ".join(str(q) for q in quantiles)),
        (
            "Input",
            "Float32 normalized RGB model inputs; denormalize in float64 using ImageNet mean/std without clipping.",
        ),
        (
            "FFT",
            "Average RGB channels; subtract each image's spatial mean; numpy.fft.fft2 with norm=forward.",
        ),
        (
            "RadialBins",
            "Original BA bins: floor(radius * 256 + 1e-10), rings 1..127; exclude DC and radii >= 0.5 cycles/pixel.",
        ),
        (
            "Amplitude",
            "Mean absolute Fourier magnitude per coefficient in each radial ring. Matches the single-sample BA analysis.",
        ),
        (
            "Power",
            "Mean squared Fourier magnitude per coefficient in each radial ring. Matches the appendix BA analysis.",
        ),
        (
            "RelativePowerPercent",
            "If a reference is supplied: 100 * (sample ring power / corresponding image reference ring power - 1).",
        ),
        (
            "Aggregation",
            "Compute spectra for individual samples first, then means and quantiles over all selected samples at each frequency. Each sample has equal weight; indices with more samples have more weight.",
        ),
        (
            "QuantileMethod",
            "numpy.quantile, method=linear. These are sample-distribution percentiles, not confidence intervals. A single sample has identical mean and quantiles.",
        ),
        (
            "Coverage",
            "Every observed numeric index must be complete and valid. Missing selected indices/samples fail. The staging group _pending is excluded.",
        ),
        (
            "SamplePositions",
            "Zero-based positions in each index's samples dataset; unrelated to SSIM representative ranks.",
        ),
    ):
        definitions.append(row)
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    selection = (
        "all" if args.index is None else f"index{args.index}_sample{args.sample_index}"
    )
    output = (
        args.output
        or base
        / "Stats"
        / "SampleFourierAnalysis"
        / f"fourier_{args.sampler}_{selection}.xlsx"
    )
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must have the .xlsx extension.")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Saved {output}: {len(coverage)} images, {sample_count} samples", flush=True)
    return output


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--sampler",
        required=True,
        help="Exact sampler name without samples_ prefix or .h5 suffix.",
    )
    parser.add_argument(
        "--index", type=int, help="Image index; requires --sample-index."
    )
    parser.add_argument(
        "--sample-index", type=int, help="Zero-based sample position; requires --index."
    )
    parser.add_argument("--quantiles", type=float, nargs="+", default=[0.1, 0.9])
    parser.add_argument(
        "--reference-sampler",
        help="Optional original/Vanilla sampler with one stored sample per image.",
    )
    parser.add_argument("--results-dir", type=Path, default=RESULTS_ROOT)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--batch-size", type=int, default=16)
    return parser.parse_args()


if __name__ == "__main__":
    analyze(parse_args())
