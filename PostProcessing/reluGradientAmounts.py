"""Measure absolute logit-gradient mass before/after each ReLU backward gate."""

import argparse
import hashlib
import json
import os
from pathlib import Path
import sys
import time

import h5py
import numpy as np
import torch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from Helper.model import load_vgg16  # noqa: E402
from Helper.paths import RESULTS_ROOT  # noqa: E402

DEFAULT_SAMPLERS = [
    "AllSamples_VanillaGradient_n1",
    "AllSamples_ADM_n100_strength0.08",
    "AllSamples_SmoothGrad_n100_map_to_sphereFalse_per_channelFalse_percentage0.09_sampling_modenormal",
]
METRICS = ("incoming_l1", "passed_l1", "blocked_fraction")


def measure(model, samples, label, device):
    """Return per-sample scalar metrics, never persist full gradients."""
    modules = [(n, m) for n, m in model.named_modules() if isinstance(m, torch.nn.ReLU)]
    if not modules:
        raise ValueError("Model contains no ReLUs")
    old_inplace = [m.inplace for _, m in modules]
    activations = {}
    handles = []
    values = {
        name: np.empty((len(samples), 3), dtype=np.float64) for name, _ in modules
    }
    values["INPUT"] = np.full((len(samples), 3), np.nan, dtype=np.float64)
    shapes = {}
    for name, module in modules:
        module.inplace = False
        handles.append(
            module.register_forward_hook(
                lambda m, i, o, name=name: activations.__setitem__(name, o)
            )
        )
    try:
        for row, sample in enumerate(samples):
            activations.clear()
            x = torch.as_tensor(
                np.asarray(sample, dtype=np.float32), device=device
            ).unsqueeze(0)
            x.requires_grad_(True)
            score = model(x)[0, label]
            outputs = [activations[name] for name, _ in modules]
            gradients = torch.autograd.grad(score, outputs + [x])
            for (name, _), output, gradient in zip(modules, outputs, gradients[:-1]):
                magnitude = gradient.detach().abs()
                incoming = magnitude.sum(dtype=torch.float64)
                passed = (magnitude * (output.detach() > 0)).sum(dtype=torch.float64)
                incoming, passed = incoming.item(), passed.item()
                if not np.isfinite([incoming, passed]).all():
                    raise ValueError(f"Non-finite gradient at {name}, sample {row}")
                values[name][row] = (
                    incoming,
                    passed,
                    (1 - passed / incoming if incoming > 0 else np.nan),
                )
                shapes[name] = tuple(output.shape[1:])
            input_l1 = gradients[-1].detach().abs().sum(dtype=torch.float64).item()
            if not np.isfinite(input_l1):
                raise ValueError("Non-finite input gradient")
            values["INPUT"][row, 0] = input_l1
            shapes["INPUT"] = tuple(x.shape[1:])
            if (row + 1) % 25 == 0:
                print(f"    {row + 1}/{len(samples)} samples", flush=True)
            del gradients, outputs, score, x
            activations.clear()
    finally:
        for handle in handles:
            handle.remove()
        for (_, module), inplace in zip(modules, old_inplace):
            module.inplace = inplace
    return values, shapes


def finite_mean(values):
    valid = np.asarray(values)[np.isfinite(values)]
    return float(valid.mean()) if len(valid) else np.nan


def interval(values, repetitions, seed):
    values = np.asarray(values, dtype=np.float64)
    values = values[np.isfinite(values)]
    if not len(values):
        return np.nan, np.nan, np.nan
    rng = np.random.default_rng(seed)
    boot = np.empty(repetitions)
    for start in range(0, repetitions, 256):
        size = min(256, repetitions - start)
        boot[start : start + size] = values[
            rng.integers(len(values), size=(size, len(values)))
        ].mean(1)
    return float(values.mean()), *np.percentile(boot, [2.5, 97.5]).tolist()


DEFINITIONS = [
    (
        "Target",
        "Stored ground-truth class logit; same label for all samplers at an index; model.eval().",
    ),
    (
        "incoming_l1",
        "Sum(abs(d logit / d h)), h = post-ReLU activation. Incoming gradient mass in the BACKWARD pass, before the ReLU mask. All channels and positions retained in the sum.",
    ),
    (
        "passed_l1",
        "Sum(abs(d logit / d h) * (h > 0)). Equals L1 norm of gradient with respect to ReLU preactivation z. Gradient mass after masking in the BACKWARD pass. PyTorch uses zero derivative at z=0.",
    ),
    (
        "blocked_fraction",
        "1 - passed_l1 / incoming_l1 per sample. Fraction of incoming L1 gradient mass blocked by this ReLU, in [0,1]. Undefined if incoming_l1=0. This weights gates by their gradient magnitude, unlike the inactive-unit fraction.",
    ),
    (
        "INPUT / incoming_l1",
        "L1 norm of the gradient with respect to the normalized model input. This is an individual-sample gradient, not the final averaged SmoothGrad attribution. No gate or passed/blocked values exist for INPUT.",
    ),
    (
        "Layer interpretation",
        "Each row refers to one ReLU module, not a whole convolutional block. Layer names are in network order in Layers. INPUT is separately reported.",
    ),
    (
        "Aggregation",
        "Mean of finite per-sample measurements within each image, then equal-weight mean of image means. Mean blocked_fraction is the mean of sample ratios, not a ratio of mean masses. No across-layer aggregation is performed.",
    ),
    (
        "Bootstrap",
        "95% percentile bootstrap over image means, with replacement, 10000 repetitions by default. Samples and pairs are not independent bootstrap units. Seed and repetition count recorded below. One-image intervals are degenerate.",
    ),
    (
        "ValidValues / ValidIndices",
        "Number of finite sample measurements / image means. NaN values are excluded and appear blank in Excel; zero incoming mass is still a valid mass but has undefined blocked_fraction.",
    ),
    (
        "Details",
        "details/index_N.h5: samplers/<integer>/layers/<module> is an [samples,3] array, columns incoming_l1, passed_l1, blocked_fraction; sample row is its zero-based index in the source file. Sampler group name attribute identifies sampler; shape attribute identifies layer geometry.",
    ),
    (
        "Reproducibility",
        "Source sample bytes, target labels, model state, software versions and script are fingerprinted. Atomic completed image files are reused only for identical fingerprints. Old analysis results are untouched; no samplers run.",
    ),
    (
        "Limits",
        "This directly measures local ReLU gradient masking. It does not isolate the causal effect of noise: incoming gradients and downstream states also change. L1 masses from different layers use different coordinates, scales and dimensions; compare conditions within a layer.",
    ),
]


def excel(path, details, config):
    book = Workbook()
    overview = book.active
    overview.title = "Overview"
    overview.append(("Entry", "Definition"))
    for row in DEFINITIONS:
        overview.append(row)
    for key, value in config.items():
        overview.append((key, json.dumps(value, default=str)))
    layers = book.create_sheet("Layers")
    layers.append(("Layer", "Shape", "Entries"))
    per_index = book.create_sheet("PerIndex")
    per_index.append(
        ("Sampler", "Index", "Layer", "Metric", "Mean", "ValidValues", "TotalValues")
    )
    summary = book.create_sheet("Summary")
    summary.append(
        (
            "Sampler",
            "Layer",
            "Metric",
            "Mean",
            "CI 2.5%",
            "CI 97.5%",
            "ValidIndices",
            "TotalIndices",
        )
    )
    coverage = book.create_sheet("Coverage")
    coverage.append(("Sampler", "Index", "SampleCount", "Status"))
    groups, seen = {}, set()
    for index, file in details:
        with h5py.File(file) as f:
            for sampler in f["samplers"].values():
                name = sampler.attrs["name"]
                coverage.append((name, index, int(sampler.attrs["count"]), "complete"))
                for layer, dataset in sampler["layers"].items():
                    shape = tuple(int(i) for i in dataset.attrs["shape"])
                    if (layer, shape) not in seen:
                        layers.append((layer, str(shape), int(np.prod(shape))))
                        seen.add((layer, shape))
                    data = dataset[:]
                    for column, metric in enumerate(METRICS):
                        if layer == "INPUT" and column != 0:
                            continue
                        value = finite_mean(data[:, column])
                        valid = int(np.isfinite(data[:, column]).sum())
                        per_index.append(
                            (name, index, layer, metric, clean(value), valid, len(data))
                        )
                        groups.setdefault((name, layer, metric), []).append(value)
    for key, values in groups.items():
        key_seed = int.from_bytes(
            hashlib.sha256(repr(key).encode()).digest()[:8], "little"
        )
        avg, low, high = interval(
            values, config["bootstrap_repetitions"], [config["seed"], key_seed]
        )
        summary.append(
            (
                *key,
                clean(avg),
                clean(low),
                clean(high),
                int(np.isfinite(values).sum()),
                len(values),
            )
        )
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 24
        sheet.column_dimensions["A"].width = 55
    overview.column_dimensions["B"].width = 110
    for row in overview.iter_rows(min_row=2):
        overview.row_dimensions[row[0].row].height = 60
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    temp = path.with_suffix(".tmp.xlsx")
    book.save(temp)
    os.replace(temp, path)


def clean(value):
    return float(value) if np.isfinite(value) else None


def run(args, model=None):
    if args.bootstrap_repetitions < 1:
        raise ValueError("bootstrap repetitions must be positive")
    names = list(dict.fromkeys(args.samplers))
    root = Path(args.results_dir) / args.model_name / args.dataset_name
    paths = [root / "Samples/allSamples" / f"samples_{name}.h5" for name in names]
    output = Path(args.output_dir) if args.output_dir else root / "ReLUGradientAmounts"
    output.mkdir(parents=True, exist_ok=True)
    (output / "details").mkdir(exist_ok=True)
    handles = []
    try:
        for path in paths:
            handles.append(h5py.File(path, "r"))
        sets = []
        for name, handle in zip(names, handles):
            for key, expected in (
                ("model_name", args.model_name),
                ("dataset_name", args.dataset_name),
            ):
                if handle.attrs.get(key) != expected:
                    raise ValueError(f"{name}: {key} mismatch")
            sets.append(
                {
                    int(k)
                    for k in handle
                    if k.isdigit() and handle[k].attrs.get("complete", False)
                }
            )
        indices = (
            sorted(set(args.indices))
            if args.indices is not None
            else sorted(set.union(*sets))
        )
        if not indices:
            raise ValueError("No complete indices")
        for name, complete in zip(names, sets):
            if set(indices) - complete:
                raise ValueError(
                    f"{name}: missing complete indices {sorted(set(indices) - complete)}"
                )
        # Validate all selected groups before starting expensive work.
        for index in indices:
            groups = [h[str(index)] for h in handles]
            labels = {int(g.attrs["label"]) for g in groups}
            shapes = {g["samples"].shape[1:] for g in groups}
            if len(labels) != 1 or len(shapes) != 1:
                raise ValueError(f"Label/geometry mismatch at {index}")
            for name, group in zip(names, groups):
                if len(group["samples"]) < 1:
                    raise ValueError(f"{name}/{index}: no samples")
        device = (
            "cuda"
            if args.device == "auto" and torch.cuda.is_available()
            else "cpu"
            if args.device == "auto"
            else args.device
        )
        model = model if model is not None else load_vgg16(device)
        model.to(device).eval()
        base = hashlib.sha256(Path(__file__).read_bytes())
        base.update(str(torch.__version__).encode())
        base.update(np.__version__.encode())
        for name, tensor in model.state_dict().items():
            base.update(name.encode())
            base.update(tensor.detach().cpu().contiguous().numpy().tobytes())
        config = dict(
            vars(args),
            resolved_device=device,
            run_sha256=base.hexdigest(),
            torch_version=str(torch.__version__),
        )
        details = []
        for index in indices:
            start = time.monotonic()
            digest = base.copy()
            for name, handle in zip(names, handles):
                group = handle[str(index)]
                digest.update(
                    json.dumps(
                        (name, int(group.attrs["label"]), group["samples"].shape)
                    ).encode()
                )
                for sample in group["samples"]:
                    digest.update(np.asarray(sample, dtype=np.float32).tobytes())
            fingerprint = digest.hexdigest()
            path = output / "details" / f"index_{index}.h5"
            reuse = False
            if path.exists() and not args.force:
                with h5py.File(path) as existing:
                    reuse = (
                        bool(existing.attrs.get("complete", False))
                        and existing.attrs.get("fingerprint") == fingerprint
                    )
            if not reuse:
                temporary = path.with_suffix(".tmp.h5")
                with h5py.File(temporary, "w", track_order=True) as result:
                    result.attrs["fingerprint"] = fingerprint
                    result.attrs["config"] = json.dumps(config, default=str)
                    samplers = result.create_group("samplers", track_order=True)
                    for number, (name, handle) in enumerate(zip(names, handles)):
                        group = handle[str(index)]
                        print(f"index {index}: {name}", flush=True)
                        values, shapes = measure(
                            model, group["samples"], int(group.attrs["label"]), device
                        )
                        target = samplers.create_group(str(number))
                        target.attrs["name"] = name
                        target.attrs["count"] = len(group["samples"])
                        target.attrs["label"] = int(group.attrs["label"])
                        layers = target.create_group("layers", track_order=True)
                        for layer, data in values.items():
                            ds = layers.create_dataset(layer, data=data)
                            ds.attrs["shape"] = shapes[layer]
                            ds.attrs["columns"] = json.dumps(METRICS)
                    result.attrs["complete"] = True
                os.replace(temporary, path)
            print(
                f"index {index}: {'reused' if reuse else 'saved'}; {time.monotonic() - start:.1f}s",
                flush=True,
            )
            details.append((index, path))
        excel(output / "relu_gradient_amounts.xlsx", details, config)
        print(f"Saved {output / 'relu_gradient_amounts.xlsx'}", flush=True)
        return output / "relu_gradient_amounts.xlsx"
    finally:
        for handle in handles:
            handle.close()


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--samplers", nargs="+", default=DEFAULT_SAMPLERS)
    parser.set_defaults(model_name="VGG16", dataset_name="IMAGENET_256")
    parser.add_argument("--results-dir", default=str(RESULTS_ROOT))
    parser.add_argument("--output-dir")
    parser.add_argument("--indices", nargs="+", type=int)
    parser.add_argument("--device", choices=("auto", "cpu", "cuda"), default="auto")
    parser.add_argument("--bootstrap-repetitions", type=int, default=10000)
    parser.add_argument("--seed", type=int, default=0)
    parser.add_argument("--force", action="store_true")
    return parser.parse_args()


if __name__ == "__main__":
    run(parse_args())
