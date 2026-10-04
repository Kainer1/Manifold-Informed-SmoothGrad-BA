"""Forward-only ReLU counts for every saved AllSamples sampler."""

import argparse
import hashlib
import json
import os
from pathlib import Path
import sys

import h5py
import numpy as np
import torch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from Helper.model import load_vgg16  # noqa: E402
from Helper.paths import RESULTS_ROOT  # noqa: E402
from PostProcessing.reluGradientAmounts import interval  # noqa: E402


def count_active(model, samples, device, batch_size=8):
    if batch_size < 1:
        raise ValueError("batch size must be positive")
    modules = [(n, m) for n, m in model.named_modules() if isinstance(m, torch.nn.ReLU)]
    if not modules:
        raise ValueError("No ReLU modules found")
    names = [n for n, _ in modules]
    counts = np.empty((len(samples), len(names)), dtype=np.int64)
    current, shapes = {}, {}

    def hook(name, output):
        if name in current:
            raise ValueError(
                f"Reused ReLU module {name}: expected one invocation per forward"
            )
        shapes[name] = tuple(output.shape[1:])
        current[name] = (output > 0).flatten(1).sum(1).cpu().numpy()

    handles = [
        m.register_forward_hook(lambda m, i, o, n=n: hook(n, o)) for n, m in modules
    ]
    try:
        model.eval()
        with torch.inference_mode():
            for start in range(0, len(samples), batch_size):
                current.clear()
                x = torch.as_tensor(
                    np.asarray(samples[start : start + batch_size], dtype=np.float32),
                    device=device,
                )
                model(x)
                for column, name in enumerate(names):
                    counts[start : start + len(x), column] = current[name]
    finally:
        for handle in handles:
            handle.remove()
    return names, shapes, counts


def observations(names, units, counts):
    for col, name in enumerate(names):
        yield name, "active_count", counts[:, col]
        yield name, "active_fraction", counts[:, col] / units[col]
    yield "ALL_UNITS", "active_count", counts.sum(1)
    yield "ALL_UNITS", "active_fraction", counts.sum(1) / units.sum()
    yield (
        "ALL_LAYERS_EQUAL_WEIGHT",
        "active_fraction",
        (counts / units[None, :]).mean(1),
    )


DEFINITIONS = [
    (
        "Purpose",
        "Forward-pass activations only. No gradients, pairwise comparisons, or new sampling.",
    ),
    (
        "active_count",
        "Number of strictly positive ReLU outputs in one layer for a sample. Each channel/location is one unit. Zero is inactive.",
    ),
    (
        "active_fraction",
        "active_count / number of units in that layer. Values in [0,1]; 0.5 means 50%.",
    ),
    (
        "ALL_UNITS",
        "Sum active counts across layers; fraction = total active count / total units. Large layers therefore contribute more.",
    ),
    (
        "ALL_LAYERS_EQUAL_WEIGHT",
        "Arithmetic mean of layer active fractions. Every layer has equal weight; no active_count reported for this statistic.",
    ),
    (
        "Aggregation",
        "First average over samples within each image index, then give each image mean equal weight. Original/VanillaGradient has one sample per image.",
    ),
    (
        "95% CI",
        "Percentile bootstrap of image means with replacement, 10000 repeats by default. Not a percentile interval over individual samples. One-image CI is degenerate.",
    ),
    (
        "Coverage",
        "Only complete, nonempty index groups are analyzed. Missing, pending or incomplete requested/observed indices are explicitly recorded. Samplers can have different coverage; consult counts.",
    ),
    (
        "Details",
        "details/<sampler>.h5: index/counts has shape [samples,layers] and int64 active counts. Index attrs layer_names, units, shapes record column identity. Row matches source sample number. Fractions are exactly recoverable from counts/units.",
    ),
    (
        "Resume",
        "Completed index results reused only when source bytes, code, weights, model structure, software, device, and batch size fingerprint match. Incomplete output groups are recomputed.",
    ),
    (
        "Interpretation",
        "Forward activity is independent of target class and does not directly quantify gradient flow. Counts alone do not establish which units remain active across samples.",
    ),
]


def write_excel(path, records, coverage, config):
    book = Workbook()
    overview = book.active
    overview.title = "Overview"
    overview.append(("Entry", "Definition"))
    for row in DEFINITIONS:
        overview.append(row)
    for k, v in config.items():
        overview.append((k, json.dumps(v, default=str)))
    layers = book.create_sheet("Layers")
    layers.append(("Layer", "Shape", "Units"))
    per = book.create_sheet("PerIndex")
    per.append(("Sampler", "Index", "Layer", "Metric", "Mean", "SampleCount"))
    summary = book.create_sheet("Summary")
    summary.append(
        (
            "Sampler",
            "Layer",
            "Metric",
            "Mean",
            "CI 2.5%",
            "CI 97.5%",
            "IndexCount",
            "MinSamples",
            "MaxSamples",
        )
    )
    cov = book.create_sheet("Coverage")
    cov.append(("Sampler", "Index", "Status", "SampleCount"))
    for row in coverage:
        cov.append(row)
    combined = {}
    seen = set()
    for sampler, index, file in records:
        with h5py.File(file) as f:
            g = f[str(index)]
            names = json.loads(g.attrs["layer_names"])
            units = np.asarray(g.attrs["units"])
            shapes = json.loads(g.attrs["shapes"])
            counts = g["counts"][:]
            for name, unit in zip(names, units):
                key = (name, tuple(shapes[name]))
                if key not in seen:
                    layers.append((name, str(tuple(shapes[name])), int(unit)))
                    seen.add(key)
            for layer, metric, values in observations(names, units, counts):
                mean = float(np.mean(values))
                key = (sampler, layer, metric)
                per.append((*key[:1], index, layer, metric, mean, len(counts)))
                combined.setdefault(key, []).append((mean, len(counts)))
    for key, rows in combined.items():
        metric_seed = int.from_bytes(
            hashlib.sha256(repr(key).encode()).digest()[:8], "little"
        )
        avg, lo, hi = interval(
            [r[0] for r in rows],
            config["bootstrap_repetitions"],
            [config["seed"], metric_seed],
        )
        summary.append(
            (
                *key,
                avg,
                lo,
                hi,
                len(rows),
                min(r[1] for r in rows),
                max(r[1] for r in rows),
            )
        )
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 25
        sheet.column_dimensions["A"].width = 55
    overview.column_dimensions["B"].width = 110
    for row in overview.iter_rows(min_row=2):
        overview.row_dimensions[row[0].row].height = 60
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    tmp = path.with_suffix(".tmp.xlsx")
    book.save(tmp)
    os.replace(tmp, path)


def run(args, model=None):
    if args.batch_size < 1 or args.bootstrap_repetitions < 1:
        raise ValueError("Positive batch size and bootstrap repetitions required")
    root = Path(args.results_dir) / args.model_name / args.dataset_name
    paths = sorted((root / "Samples/allSamples").glob("samples_*.h5"))
    if not paths:
        raise ValueError("No AllSamples files found")
    output = Path(args.output_dir) if args.output_dir else root / "ReLUActiveAnalysis"
    (output / "details").mkdir(parents=True, exist_ok=True)
    device = (
        ("cuda" if torch.cuda.is_available() else "cpu")
        if args.device == "auto"
        else args.device
    )
    model = model if model is not None else load_vgg16(device)
    model.to(device).eval()
    base = hashlib.sha256(Path(__file__).read_bytes())
    base.update(repr(model).encode())
    base.update(
        json.dumps(
            (str(torch.__version__), np.__version__, device, args.batch_size)
        ).encode()
    )
    for name, tensor in model.state_dict().items():
        base.update(name.encode())
        base.update(tensor.detach().cpu().contiguous().numpy().tobytes())
    config = dict(
        vars(args),
        resolved_device=device,
        model_run_sha256=base.hexdigest(),
        source_files=[str(p) for p in paths],
    )
    records = []
    coverage = []
    for path in paths:
        name = path.stem.removeprefix("samples_")
        out = output / "details" / f"{name}.h5"
        try:
            source = h5py.File(path, "r")
        except OSError as error:
            coverage.append((name, "*", f"unreadable/locked: {error}", 0))
            continue
        with source, h5py.File(out, "a") as target:
            for key, expected in (
                ("model_name", args.model_name),
                ("dataset_name", args.dataset_name),
            ):
                if source.attrs.get(key) != expected:
                    raise ValueError(f"{path}: {key} mismatch")
            target.attrs["config"] = json.dumps(config, default=str)
            target.attrs["sampler"] = name
            keys = sorted(
                {int(k) for k in source if k.isdigit()}
                | {int(k) for k in source.get("_pending", {}) if k.isdigit()}
            )
            for index in (
                sorted(set(args.indices)) if args.indices is not None else keys
            ):
                group = source.get(str(index))
                if (
                    group is None
                    or not group.attrs.get("complete", False)
                    or "samples" not in group
                    or len(group["samples"]) == 0
                ):
                    coverage.append((name, index, "missing/pending/incomplete", 0))
                    continue
                samples = group["samples"]
                digest = base.copy()
                digest.update(json.dumps((name, index, samples.shape)).encode())
                for start in range(0, len(samples), args.batch_size):
                    digest.update(
                        np.asarray(
                            samples[start : start + args.batch_size], dtype=np.float32
                        ).tobytes()
                    )
                fingerprint = digest.hexdigest()
                key = str(index)
                reuse = (
                    key in target
                    and target[key].attrs.get("complete", False)
                    and target[key].attrs.get("fingerprint") == fingerprint
                )
                if not reuse:
                    names, shapes, counts = count_active(
                        model, samples, device, args.batch_size
                    )
                    pending = f"_pending_{index}"
                    if pending in target:
                        del target[pending]
                    dest = target.create_group(pending)
                    dest.create_dataset("counts", data=counts)
                    dest.attrs["layer_names"] = json.dumps(names)
                    dest.attrs["shapes"] = json.dumps(shapes)
                    dest.attrs["units"] = [int(np.prod(shapes[n])) for n in names]
                    dest.attrs["fingerprint"] = fingerprint
                    dest.attrs["complete"] = True
                    target.flush()
                    if key in target:
                        del target[key]
                    target.move(pending, key)
                    target.flush()
                records.append((name, index, out))
                coverage.append((name, index, "complete", len(samples)))
                print(
                    f"{name}: index {index}, {len(samples)} samples, {'reused' if reuse else 'saved'}",
                    flush=True,
                )
    write_excel(output / "relu_active.xlsx", records, coverage, config)
    print(
        f"Saved {output / 'relu_active.xlsx'}; {len(records)} sampler/image groups",
        flush=True,
    )
    return output / "relu_active.xlsx"


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results-dir", default=str(RESULTS_ROOT))
    parser.set_defaults(model_name="VGG16", dataset_name="IMAGENET_256")
    parser.add_argument("--output-dir")
    parser.add_argument("--indices", nargs="+", type=int)
    parser.add_argument("--batch-size", type=int, default=8)
    parser.add_argument("--device", choices=["auto", "cpu", "cuda"], default="auto")
    parser.add_argument("--bootstrap-repetitions", type=int, default=10000)
    parser.add_argument("--seed", type=int, default=0)
    return parser.parse_args()


if __name__ == "__main__":
    run(parse_args())
