"""Evaluate similarity among all sample attributions of one or more samplers."""

import argparse
import sys
from pathlib import Path

import h5py
import numpy as np
import torch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from scipy.fft import dctn
from scipy.stats import rankdata


CODE_DIR = Path(__file__).resolve().parents[1]
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from Helper.model import load_vgg16  # noqa: E402
from Helper.paths import RESULTS_ROOT  # noqa: E402


PAIRWISE_METRICS = (
    ("cosine", "Cosine"),
    ("relative_l1", "Relative L1 Difference"),
    ("spearman", "Spearman"),
)
COARSE_SCALES = (8, 16, 32, 64)
SAMPLE_METRICS = (
    ("complexity", "Complexity"),
    ("sparseness", "Sparseness"),
    *((f"coarse_{scale}", f"Coarse {scale}") for scale in COARSE_SCALES),
)
ALL_METRICS = PAIRWISE_METRICS + SAMPLE_METRICS


def _resolve_device(device: str) -> torch.device:
    if device == "auto":
        return torch.device("cuda" if torch.cuda.is_available() else "cpu")
    resolved = torch.device(device)
    if resolved.type == "cuda" and not torch.cuda.is_available():
        raise RuntimeError("CUDA was requested, but no CUDA device is available.")
    return resolved


def _numeric_keys(h5_file: h5py.File) -> list[str]:
    keys = [key for key in h5_file.keys() if key != "_pending"]
    try:
        return sorted(keys, key=int)
    except ValueError as error:
        raise ValueError(
            "All sample groups must be integer dataset indices."
        ) from error


def compute_l2_pooled_attributions(
    model,
    samples: h5py.Dataset,
    label: int,
    device: torch.device,
    batch_size: int,
) -> torch.Tensor:
    """Recompute raw sample gradients and return RGB-L2-pooled heatmaps."""
    pooled_batches = []
    for start in range(0, samples.shape[0], batch_size):
        stop = min(start + batch_size, samples.shape[0])
        inputs = torch.from_numpy(samples[start:stop].astype(np.float32)).to(device)
        inputs.requires_grad_(True)
        logits = model(inputs)
        selected_logits = logits[:, int(label)].sum()
        gradients = torch.autograd.grad(selected_logits, inputs)[0]
        pooled = torch.linalg.vector_norm(gradients.detach(), ord=2, dim=1)
        pooled_batches.append(pooled.cpu())
    return torch.cat(pooled_batches, dim=0)


def _pairwise_cosine(flattened: torch.Tensor) -> np.ndarray:
    norms = torch.linalg.vector_norm(flattened, ord=2, dim=1)
    valid_rows = norms > torch.finfo(flattened.dtype).eps
    normalized = flattened / norms.clamp_min(torch.finfo(flattened.dtype).eps)[:, None]
    similarities = normalized @ normalized.T
    row_indices, column_indices = torch.triu_indices(
        flattened.shape[0], flattened.shape[0], offset=1, device=flattened.device
    )
    values = similarities[row_indices, column_indices]
    valid_pairs = valid_rows[row_indices] & valid_rows[column_indices]
    values = torch.where(valid_pairs, values, torch.nan)
    return values.float().cpu().numpy()


def _pairwise_relative_l1(flattened: torch.Tensor) -> np.ndarray:
    magnitudes = flattened.abs().sum(dim=1)
    denominator = magnitudes[:, None] + magnitudes[None, :]
    distances = 2.0 * (magnitudes[:, None] - magnitudes[None, :]).abs()
    distances = distances / denominator.clamp_min(torch.finfo(flattened.dtype).eps)
    row_indices, column_indices = torch.triu_indices(
        flattened.shape[0], flattened.shape[0], offset=1, device=flattened.device
    )
    values = distances[row_indices, column_indices]
    valid_pairs = (
        denominator[row_indices, column_indices] > torch.finfo(flattened.dtype).eps
    )
    values = torch.where(valid_pairs, values, torch.nan)
    return values.float().cpu().numpy()


def _pairwise_spearman(flattened: torch.Tensor, device: torch.device) -> np.ndarray:
    ranks = rankdata(
        flattened.cpu().numpy(),
        axis=1,
        method="average",
    ).astype(np.float32, copy=False)
    rank_tensor = torch.from_numpy(ranks).to(device)
    rank_tensor -= rank_tensor.mean(dim=1, keepdim=True)
    return _pairwise_cosine(rank_tensor)


def calculate_pairwise_metrics(
    heatmaps: torch.Tensor,
    device: torch.device,
) -> dict[str, np.ndarray]:
    """Calculate every unordered sample-pair value for the three metrics."""
    if heatmaps.ndim != 3:
        raise ValueError(
            f"Expected pooled heatmaps with shape [samples, height, width], "
            f"but received {tuple(heatmaps.shape)}."
        )
    if heatmaps.shape[0] < 2:
        raise ValueError("At least two samples are required for pairwise metrics.")

    maps_on_device = heatmaps.to(device=device, dtype=torch.float32)
    flattened = maps_on_device.flatten(start_dim=1)
    return {
        "cosine": _pairwise_cosine(flattened),
        "relative_l1": _pairwise_relative_l1(flattened),
        "spearman": _pairwise_spearman(flattened, device),
    }


def calculate_sample_metrics(heatmaps: torch.Tensor) -> dict[str, np.ndarray]:
    """Calculate entropy complexity and coarse-scale energy per sample heatmap."""
    if heatmaps.ndim != 3:
        raise ValueError(
            f"Expected pooled heatmaps with shape [samples, height, width], "
            f"but received {tuple(heatmaps.shape)}."
        )

    maps = heatmaps.float().cpu().numpy()
    flattened = np.abs(maps.reshape(maps.shape[0], -1)).astype(np.float64, copy=False)
    attribution_sums = flattened.sum(axis=1, keepdims=True)
    valid_complexity = attribution_sums[:, 0] > np.finfo(np.float64).eps
    probabilities = np.divide(
        flattened,
        attribution_sums,
        out=np.zeros_like(flattened),
        where=attribution_sums > 0,
    )
    log_probabilities = np.zeros_like(probabilities)
    np.log(probabilities, out=log_probabilities, where=probabilities > 0)
    complexity = -(probabilities * log_probabilities).sum(axis=1)
    complexity[~valid_complexity] = np.nan

    centered = maps - maps.mean(axis=(-2, -1), keepdims=True)
    coefficients = dctn(centered, axes=(-2, -1), norm="ortho")
    power = np.square(coefficients.astype(np.float64, copy=False))
    power[:, 0, 0] = 0.0
    total_power = power.sum(axis=(-2, -1))
    valid_power = total_power > np.finfo(np.float64).eps

    height, width = maps.shape[-2:]
    vertical_frequency = np.arange(height, dtype=np.float64)[:, None] / height
    horizontal_frequency = np.arange(width, dtype=np.float64)[None, :] / width
    radial_frequency = np.sqrt(
        np.square(vertical_frequency) + np.square(horizontal_frequency)
    )

    sorted_attributions = np.sort(flattened, axis=1)
    n_features = flattened.shape[1]
    ranks = np.arange(1, n_features + 1, dtype=np.float64)
    gini_weights = 2.0 * ranks - n_features - 1.0
    gini_numerator = (sorted_attributions * gini_weights).sum(axis=1)
    sparseness = np.divide(
        gini_numerator,
        n_features * attribution_sums[:, 0],
        out=np.full(maps.shape[0], np.nan, dtype=np.float64),
        where=valid_complexity,
    )

    results = {"complexity": complexity, "sparseness": sparseness}
    for scale in COARSE_SCALES:
        # A DCT basis with radial frequency f has spatial wavelength 2 / f.
        low_frequency_mask = radial_frequency <= (2.0 / scale)
        low_frequency_mask[0, 0] = False
        low_frequency_power = power[:, low_frequency_mask].sum(axis=1)
        score = np.divide(
            low_frequency_power,
            total_power,
            out=np.full_like(total_power, np.nan),
            where=valid_power,
        )
        results[f"coarse_{scale}"] = score
    return results


def _finite_mean(values: np.ndarray) -> tuple[float, int]:
    finite = values[np.isfinite(values)]
    if finite.size == 0:
        return float("nan"), 0
    return float(finite.mean(dtype=np.float64)), int(finite.size)


def bootstrap_mean_confidence_interval(
    values: list[float],
    repetitions: int,
    rng: np.random.Generator,
) -> tuple[float, float, float]:
    """Return the mean and a percentile-bootstrap 95% interval over images."""
    finite = np.asarray(values, dtype=np.float64)
    finite = finite[np.isfinite(finite)]
    if finite.size == 0:
        return float("nan"), float("nan"), float("nan")
    if repetitions <= 0:
        raise ValueError("bootstrap_repetitions must be greater than zero.")

    selections = rng.integers(0, finite.size, size=(repetitions, finite.size))
    bootstrap_means = finite[selections].mean(axis=1)
    lower, upper = np.percentile(bootstrap_means, (2.5, 97.5))
    return float(finite.mean()), float(lower), float(upper)


def _style_header(worksheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def _write_workbook(
    output_path: Path,
    summary_rows: list[dict],
    per_index_rows: list[dict],
    bootstrap_repetitions: int,
    seed: int,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_headers = ["Sampler", "Index Count", "Min Samples", "Max Samples"]
    for _, label in ALL_METRICS:
        summary_headers.extend((f"{label} Mean", f"{label} CI [2.5%, 97.5%]"))
    summary.append(summary_headers)
    for row in summary_rows:
        summary.append([row.get(header) for header in summary_headers])
    _style_header(summary)
    summary.column_dimensions["A"].width = 55
    for column in range(2, len(summary_headers) + 1):
        summary.column_dimensions[summary.cell(1, column).column_letter].width = 22
        for cell in summary.iter_cols(
            min_col=column, max_col=column, min_row=2
        ).__next__():
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"

    per_index = workbook.create_sheet("PerIndex")
    per_index_headers = ["Sampler", "Index", "Sample Count", "Pair Count"]
    for _, label in PAIRWISE_METRICS:
        per_index_headers.extend((f"{label} Mean", f"{label} Valid Pairs"))
    for _, label in SAMPLE_METRICS:
        per_index_headers.extend((f"{label} Mean", f"{label} Valid Samples"))
    per_index.append(per_index_headers)
    for row in per_index_rows:
        per_index.append([row.get(header) for header in per_index_headers])
    _style_header(per_index)
    per_index.column_dimensions["A"].width = 55
    for column in range(2, len(per_index_headers) + 1):
        per_index.column_dimensions[per_index.cell(1, column).column_letter].width = 22

    definitions = workbook.create_sheet("Definitions")
    definitions.append(("Item", "Definition"))
    definitions_rows = (
        (
            "Input",
            "Float32 normalized model inputs from saved sampler outputs; "
            "gradients target the stored ground-truth label.",
        ),
        ("Pooling", "L2 norm over the three gradient color channels."),
        ("Cosine", "Pairwise cosine similarity of flattened pooled heatmaps."),
        (
            "Relative L1 Difference",
            "2 * absolute difference of L1 norms divided by their sum.",
        ),
        (
            "Spearman",
            "Pearson correlation of pixel ranks; tied values receive average ranks.",
        ),
        (
            "Complexity",
            "Entropy of each sample's absolute fractional attribution mass, following "
            "the Quantus Complexity definition; values are averaged per index.",
        ),
        (
            "Sparseness",
            "Gini index of each sample's absolute attribution values, following "
            "the Quantus / Chalasani et al. definition; values are averaged per index.",
        ),
        (
            "Coarse 8/16/32/64",
            "Fraction of non-constant orthonormal 2D-DCT energy at spatial wavelengths "
            "of at least 8, 16, 32, or 64 pixels; values are averaged per index.",
        ),
        (
            "Image aggregation",
            "Pairwise metrics average all unordered pairs per index; Complexity and "
            "Coarse metrics average all individual samples per index. The Summary mean "
            "gives every index equal weight.",
        ),
        (
            "95% CI",
            f"Percentile bootstrap over complete dataset indices with "
            f"{bootstrap_repetitions} repetitions and seed {seed}.",
        ),
    )
    for row in definitions_rows:
        definitions.append(row)
    _style_header(definitions)
    definitions.column_dimensions["A"].width = 28
    definitions.column_dimensions["B"].width = 120

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_suffix(".tmp.xlsx")
    workbook.save(temporary_path)
    temporary_path.replace(output_path)


def analyze_sample_similarity(
    sampler_names: list[str],
    model_name: str = "VGG16",
    dataset_name: str = "IMAGENET_256",
    results_dir: Path | None = None,
    output_path: Path | None = None,
    batch_size: int = 32,
    bootstrap_repetitions: int = 10_000,
    seed: int = 0,
    device: str = "auto",
) -> Path:
    if batch_size <= 0:
        raise ValueError("batch_size must be greater than zero.")
    if not sampler_names:
        raise ValueError("At least one sampler name is required.")

    result_root = Path(results_dir) if results_dir else RESULTS_ROOT
    input_dir = result_root / model_name / dataset_name / "Samples" / "allSamples"
    resolved_device = _resolve_device(device)
    model = load_vgg16(resolved_device)
    model.eval()
    rng = np.random.default_rng(seed)
    summary_rows = []
    per_index_rows = []

    for sampler_name in sampler_names:
        sample_path = input_dir / f"samples_{sampler_name}.h5"
        if not sample_path.exists():
            raise FileNotFoundError(f"All-sample file not found: {sample_path}")

        sampler_index_rows = []
        with h5py.File(sample_path, "r") as h5_file:
            stored_model = str(h5_file.attrs.get("model_name", model_name))
            stored_dataset = str(h5_file.attrs.get("dataset_name", dataset_name))
            if stored_model != model_name or stored_dataset != dataset_name:
                raise ValueError(
                    f"Metadata mismatch in {sample_path}: expected "
                    f"{model_name}/{dataset_name}, found {stored_model}/{stored_dataset}."
                )

            for index in _numeric_keys(h5_file):
                group = h5_file[index]
                if not bool(group.attrs.get("complete", False)):
                    continue
                if "samples" not in group or "label" not in group.attrs:
                    raise KeyError(
                        f"Index {index} in {sample_path} lacks samples or its label."
                    )
                samples = group["samples"]
                heatmaps = compute_l2_pooled_attributions(
                    model=model,
                    samples=samples,
                    label=int(group.attrs["label"]),
                    device=resolved_device,
                    batch_size=batch_size,
                )
                if heatmaps.shape[0] >= 2:
                    metric_values = calculate_pairwise_metrics(
                        heatmaps, resolved_device
                    )
                else:
                    metric_values = {
                        metric_key: np.empty(0, dtype=np.float32)
                        for metric_key, _ in PAIRWISE_METRICS
                    }
                sample_metric_values = calculate_sample_metrics(heatmaps)
                pair_count = samples.shape[0] * (samples.shape[0] - 1) // 2
                row = {
                    "Sampler": sampler_name,
                    "Index": int(index),
                    "Sample Count": int(samples.shape[0]),
                    "Pair Count": int(pair_count),
                }
                for metric_key, metric_label in PAIRWISE_METRICS:
                    mean, valid_count = _finite_mean(metric_values[metric_key])
                    row[f"{metric_label} Mean"] = mean
                    row[f"{metric_label} Valid Pairs"] = valid_count
                for metric_key, metric_label in SAMPLE_METRICS:
                    mean, valid_count = _finite_mean(sample_metric_values[metric_key])
                    row[f"{metric_label} Mean"] = mean
                    row[f"{metric_label} Valid Samples"] = valid_count
                sampler_index_rows.append(row)
                per_index_rows.append(row)
                print(
                    f"[{sampler_name}] index {index}: "
                    f"{samples.shape[0]} samples, {pair_count} pairs",
                    flush=True,
                )

        if not sampler_index_rows:
            raise ValueError(f"No complete indices found in {sample_path}.")

        sample_counts = [row["Sample Count"] for row in sampler_index_rows]
        summary_row = {
            "Sampler": sampler_name,
            "Index Count": len(sampler_index_rows),
            "Min Samples": min(sample_counts),
            "Max Samples": max(sample_counts),
        }
        for _, metric_label in ALL_METRICS:
            image_means = [row[f"{metric_label} Mean"] for row in sampler_index_rows]
            mean, lower, upper = bootstrap_mean_confidence_interval(
                image_means, bootstrap_repetitions, rng
            )
            summary_row[f"{metric_label} Mean"] = mean
            summary_row[f"{metric_label} CI [2.5%, 97.5%]"] = (
                f"[{lower:.4f}, {upper:.4f}]"
                if np.isfinite(lower) and np.isfinite(upper)
                else "NaN"
            )
        summary_rows.append(summary_row)

    if output_path is None:
        output_path = (
            result_root
            / model_name
            / dataset_name
            / "Stats"
            / "AttributionSampleSimilarity"
            / "attribution_sample_similarity.xlsx"
        )
    output_path = Path(output_path)
    _write_workbook(
        output_path,
        summary_rows,
        per_index_rows,
        bootstrap_repetitions,
        seed,
    )
    print(f"Attribution sample similarity saved to {output_path}")
    return output_path


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Recompute attributions from all stored sampler outputs and analyze "
            "their pairwise similarity."
        )
    )
    parser.add_argument("--sampler_names", nargs="+", required=True)
    parser.set_defaults(model_name="VGG16", dataset_name="IMAGENET_256")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--batch_size", type=int, default=32)
    parser.add_argument("--bootstrap_repetitions", type=int, default=10_000)
    parser.add_argument("--seed", type=int, default=0)
    parser.add_argument("--device", default="auto")
    return parser.parse_args()


if __name__ == "__main__":
    arguments = _parse_args()
    analyze_sample_similarity(
        sampler_names=arguments.sampler_names,
        model_name=arguments.model_name,
        dataset_name=arguments.dataset_name,
        output_path=arguments.output,
        batch_size=arguments.batch_size,
        bootstrap_repetitions=arguments.bootstrap_repetitions,
        seed=arguments.seed,
        device=arguments.device,
    )
