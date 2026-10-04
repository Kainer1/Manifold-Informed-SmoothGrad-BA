"""Original gradient statistics and Excel output for the fixed thesis setup."""

from dataclasses import dataclass
from numbers import Real
from pathlib import Path

import numpy as np
import torch
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from Helper.paths import RESULTS_DIR


RAW_GRADIENT_METRICS = (
    "AllSamplesAbsSum",
    "MeanSampleAbsSum",
    "FinalAttributionAbsSum",
    "FirstSampleAbsSum",
    "SampleCount",
)

METRIC_DEFINITIONS = (
    (
        "AllSamplesAbsSum",
        "Summe von |g| über alle erzeugten Samples, Kanäle und Pixel eines "
        "Datensatzindex.",
    ),
    (
        "MeanSampleAbsSum",
        "AllSamplesAbsSum geteilt durch SampleCount; mittlere Betragssumme "
        "eines einzelnen Samples.",
    ),
    (
        "FinalAttributionAbsSum",
        "Betragssumme nach dem vorzeichenbehafteten Mitteln der Samplegradienten; "
        "entspricht der Stärke der resultierenden Attribution.",
    ),
    (
        "FirstSampleAbsSum",
        "Betragssumme des ersten vom Sampler ausgegebenen Samples.",
    ),
    ("SampleCount", "Anzahl der tatsächlich ausgewerteten Samples."),
    (
        "FirstSampleGrid",
        "16x16-Matrix des mittleren |g| im jeweiligen Block des ersten Samples "
        "und ersten angeforderten Datensatzindex. Die Rohwerte werden zweimal "
        "angezeigt: mit bildinterner und mit gemeinsamer Farbskala über alle "
        "Sampler in der Tabelle.",
    ),
    (
        "Gradientenbezug",
        "Ableitung des ausgewählten Ground-Truth-Klassenlogits bezüglich des "
        "normalisierten Modelleingangs.",
    ),
    (
        "LowGradient-Hinweis",
        "Erfasst werden nur die Gradienten aus compute_attributions. Zusätzliche "
        "Gradienten zur Konstruktion einer Low-/High-Gradient-Maske werden nicht "
        "mitgezählt.",
    ),
)


def mean_absolute_grid(gradient: torch.Tensor, grid_size: int = 16) -> torch.Tensor:
    """Average ``abs(gradient)`` over channels and equally sized grid blocks."""
    if gradient.ndim != 3:
        raise ValueError(
            "The first-sample gradient must have shape [channels, height, width], "
            f"but received {tuple(gradient.shape)}."
        )

    channels, height, width = gradient.shape
    if height % grid_size != 0 or width % grid_size != 0:
        raise ValueError(
            f"Gradient size {height}x{width} cannot be split into an exact "
            f"{grid_size}x{grid_size} grid."
        )

    block_height = height // grid_size
    block_width = width // grid_size
    blocks = (
        gradient.detach()
        .abs()
        .reshape(
            channels,
            grid_size,
            block_height,
            grid_size,
            block_width,
        )
    )
    return blocks.mean(dim=(0, 2, 4)).to(torch.float64)


@dataclass
class FirstSampleGrid:
    sampler_name: str
    dataset_index: str
    gradient_shape: tuple[int, int, int]
    values: np.ndarray

    @property
    def block_shape(self) -> tuple[int, int, int]:
        channels, height, width = self.gradient_shape
        grid_height, grid_width = self.values.shape
        return channels, height // grid_height, width // grid_width


class RawGradientAccumulator:
    """Collect per-index absolute-gradient statistics across processing batches."""

    def __init__(self, requested_indices: list, grid_size: int = 16):
        self.requested_indices = [str(index) for index in requested_indices]
        self.first_index = self.requested_indices[0] if self.requested_indices else None
        self.grid_size = grid_size
        self.metric_values = {metric: {} for metric in RAW_GRADIENT_METRICS}
        self.first_sample_grid: np.ndarray | None = None
        self.first_sample_gradient_shape: tuple[int, int, int] | None = None

    @property
    def has_data(self) -> bool:
        return bool(self.metric_values["SampleCount"])

    def grid_position(self, batch_indices: list) -> int | None:
        """Return the batch position of the first requested dataset index."""
        if self.first_index is None or self.first_sample_grid is not None:
            return None
        normalized_indices = [str(index) for index in batch_indices]
        try:
            return normalized_indices.index(self.first_index)
        except ValueError:
            return None

    def record_batch(
        self,
        batch_indices: list,
        all_samples_abs_sum: torch.Tensor,
        first_sample_abs_sum: torch.Tensor,
        final_attribution_abs_sum: torch.Tensor,
        sample_count: int,
        first_sample_grid: torch.Tensor | None = None,
        gradient_shape: tuple[int, int, int] | None = None,
    ):
        batch_indices = [str(index) for index in batch_indices]
        expected = len(batch_indices)
        tensors = {
            "AllSamplesAbsSum": all_samples_abs_sum,
            "FirstSampleAbsSum": first_sample_abs_sum,
            "FinalAttributionAbsSum": final_attribution_abs_sum,
        }
        for metric_name, values in tensors.items():
            if values.numel() != expected:
                raise ValueError(
                    f"{metric_name} contains {values.numel()} values for "
                    f"{expected} dataset indices."
                )

        all_values = all_samples_abs_sum.detach().cpu().double().flatten().tolist()
        first_values = first_sample_abs_sum.detach().cpu().double().flatten().tolist()
        final_values = (
            final_attribution_abs_sum.detach().cpu().double().flatten().tolist()
        )

        for position, index in enumerate(batch_indices):
            total = float(all_values[position])
            self.metric_values["AllSamplesAbsSum"][index] = total
            self.metric_values["MeanSampleAbsSum"][index] = total / sample_count
            self.metric_values["FinalAttributionAbsSum"][index] = float(
                final_values[position]
            )
            self.metric_values["FirstSampleAbsSum"][index] = float(
                first_values[position]
            )
            self.metric_values["SampleCount"][index] = int(sample_count)

        if first_sample_grid is not None:
            if gradient_shape is None:
                raise ValueError(
                    f"gradient_shape is required for the {self.grid_size}x"
                    f"{self.grid_size} grid."
                )
            grid_array = first_sample_grid.detach().cpu().double().numpy()
            expected_grid_shape = (self.grid_size, self.grid_size)
            if grid_array.shape != expected_grid_shape:
                raise ValueError(
                    f"Expected grid shape {expected_grid_shape}, got {grid_array.shape}."
                )
            normalized_shape = tuple(int(v) for v in gradient_shape)
            if len(normalized_shape) != 3:
                raise ValueError(
                    f"Expected a three-dimensional gradient shape, got {gradient_shape}."
                )
            self.first_sample_grid = grid_array
            self.first_sample_gradient_shape = (
                normalized_shape[0],
                normalized_shape[1],
                normalized_shape[2],
            )


def _is_index_column(column_name: str) -> bool:
    try:
        int(column_name)
    except (TypeError, ValueError):
        return False
    return True


def _read_summary_rows(workbook) -> list[dict]:
    if "Summary" not in workbook.sheetnames:
        return []

    worksheet = workbook["Summary"]
    headers = [
        str(cell.value) if cell.value is not None else "" for cell in worksheet[1]
    ]
    rows = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row = {header: value for header, value in zip(headers, values) if header}
        if row.get("Samplername") and row.get("Metrik"):
            rows.append(row)
    return rows


def _summary_rows_for_sampler(
    sampler_name: str, accumulator: RawGradientAccumulator
) -> list[dict]:
    rows = []
    for metric_name in RAW_GRADIENT_METRICS:
        values = accumulator.metric_values[metric_name]
        rows.append(
            {
                "Samplername": sampler_name,
                "Metrik": metric_name,
                **{index: values.get(index) for index in accumulator.requested_indices},
            }
        )
    return rows


def _write_summary(workbook, rows: list[dict]):
    if "Summary" in workbook.sheetnames:
        workbook.remove(workbook["Summary"])
    worksheet = workbook.create_sheet("Summary", 0)

    index_columns = sorted(
        {
            str(column)
            for row in rows
            for column in row
            if _is_index_column(str(column))
        },
        key=int,
    )
    columns = ["Samplername", "Metrik", "Mittelwert", "Gesamtsumme"] + index_columns
    worksheet.append(columns)

    metric_order = {
        name: position for position, name in enumerate(RAW_GRADIENT_METRICS)
    }
    rows.sort(
        key=lambda row: (
            str(row["Samplername"]),
            metric_order.get(str(row["Metrik"]), len(metric_order)),
        )
    )

    for row in rows:
        numeric_values = [
            float(row[column])
            for column in index_columns
            if isinstance(row.get(column), Real)
        ]
        mean = sum(numeric_values) / len(numeric_values) if numeric_values else None
        total = sum(numeric_values) if numeric_values else None
        worksheet.append(
            [row.get("Samplername"), row.get("Metrik"), mean, total]
            + [row.get(column) for column in index_columns]
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    worksheet.freeze_panes = "E2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 60
    worksheet.column_dimensions["B"].width = 28
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 18
    for column in range(5, len(columns) + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 15

    for row in range(2, worksheet.max_row + 1):
        is_count = worksheet.cell(row, 2).value == "SampleCount"
        for column in range(3, worksheet.max_column + 1):
            if is_count:
                worksheet.cell(row, column).number_format = (
                    "0.00" if column == 3 else "0"
                )
            else:
                worksheet.cell(row, column).number_format = "0.000000E+00"


def _read_grid_blocks(workbook) -> dict[str, FirstSampleGrid]:
    if "FirstSampleGrid" not in workbook.sheetnames:
        return {}

    worksheet = workbook["FirstSampleGrid"]
    blocks = {}
    row = 1
    while row <= worksheet.max_row:
        if worksheet.cell(row, 1).value != "Samplername":
            row += 1
            continue

        sampler_name = worksheet.cell(row, 2).value
        dataset_index = worksheet.cell(row + 1, 2).value
        gradient_shape_value = str(worksheet.cell(row + 2, 2).value)
        header_row = row + 5
        grid_size = 0
        while worksheet.cell(header_row, 2 + grid_size).value == grid_size:
            grid_size += 1
        try:
            if grid_size == 0:
                raise ValueError("No grid columns found.")
            gradient_shape_values = tuple(
                int(v) for v in gradient_shape_value.split("x")
            )
            values = np.array(
                [
                    [
                        worksheet.cell(row + 6 + grid_row, 2 + grid_column).value
                        for grid_column in range(grid_size)
                    ]
                    for grid_row in range(grid_size)
                ],
                dtype=float,
            )
        except (TypeError, ValueError):
            row += 1
            continue

        if sampler_name and len(gradient_shape_values) == 3:
            gradient_shape = (
                gradient_shape_values[0],
                gradient_shape_values[1],
                gradient_shape_values[2],
            )
            blocks[str(sampler_name)] = FirstSampleGrid(
                sampler_name=str(sampler_name),
                dataset_index=str(dataset_index),
                gradient_shape=gradient_shape,
                values=values,
            )
        row += grid_size + 7
    return blocks


def _color_scale_limits(values) -> tuple[float, float, float]:
    finite_values = np.asarray(values, dtype=float)
    finite_values = finite_values[np.isfinite(finite_values)]
    if finite_values.size == 0:
        return 0.0, 0.5, 1.0
    minimum = float(finite_values.min())
    maximum = float(finite_values.max())
    if maximum == minimum:
        maximum = minimum + max(abs(minimum) * 1e-12, 1e-30)
    return minimum, (minimum + maximum) / 2.0, maximum


def _add_color_scale(worksheet, cell_range: str, limits: tuple[float, float, float]):
    worksheet.conditional_formatting.add(
        cell_range,
        ColorScaleRule(
            start_type="num",
            start_value=limits[0],
            start_color="FFFFFF",
            mid_type="num",
            mid_value=limits[1],
            mid_color="FFD966",
            end_type="num",
            end_value=limits[2],
            end_color="C00000",
        ),
    )


def _write_grid_blocks(workbook, blocks: dict[str, FirstSampleGrid]):
    if "FirstSampleGrid" in workbook.sheetnames:
        workbook.remove(workbook["FirstSampleGrid"])
    worksheet = workbook.create_sheet("FirstSampleGrid", 1)
    worksheet.freeze_panes = "B7"
    maximum_grid_size = max(
        (int(block.values.shape[1]) for block in blocks.values()),
        default=16,
    )
    gap_column = maximum_grid_size + 2
    global_label_column = gap_column + 1
    global_data_column = global_label_column + 1
    last_column = global_data_column + maximum_grid_size - 1

    worksheet.column_dimensions["A"].width = 26
    worksheet.column_dimensions[get_column_letter(global_label_column)].width = 26
    worksheet.column_dimensions[get_column_letter(gap_column)].width = 3
    for column in range(2, last_column + 1):
        if column in {gap_column, global_label_column}:
            continue
        worksheet.column_dimensions[get_column_letter(column)].width = 16

    global_limits = _color_scale_limits(
        np.concatenate([block.values.ravel() for block in blocks.values()])
        if blocks
        else np.array([], dtype=float)
    )

    current_row = 1
    for sampler_name in sorted(blocks):
        block = blocks[sampler_name]
        channels, block_height, block_width = block.block_shape
        metadata = (
            ("Samplername", block.sampler_name),
            ("DatasetIndex", block.dataset_index),
            ("GradientShape", "x".join(str(v) for v in block.gradient_shape)),
            ("BlockShape", f"{channels}x{block_height}x{block_width}"),
            (
                "Definition",
                "MeanAbsoluteGradient of the first yielded sample in each block",
            ),
        )
        for offset, (label, value) in enumerate(metadata):
            worksheet.cell(current_row + offset, 1, label).font = Font(bold=True)
            worksheet.cell(current_row + offset, 2, value)

        header_row = current_row + 5
        grid_size = int(block.values.shape[0])
        if block.values.shape != (grid_size, grid_size):
            raise ValueError(
                f"Expected a square grid for {block.sampler_name}, "
                f"got {block.values.shape}."
            )
        worksheet.cell(header_row, 1, "Bildnormiert / Spalte").font = Font(bold=True)
        worksheet.cell(
            header_row,
            global_label_column,
            "Tabellennormiert / Spalte",
        ).font = Font(bold=True)
        for column in range(grid_size):
            worksheet.cell(header_row, column + 2, column).font = Font(bold=True)
            worksheet.cell(
                header_row,
                global_data_column + column,
                column,
            ).font = Font(bold=True)

        first_data_row = current_row + 6
        for grid_row in range(grid_size):
            worksheet.cell(first_data_row + grid_row, 1, grid_row).font = Font(
                bold=True
            )
            worksheet.cell(
                first_data_row + grid_row,
                global_label_column,
                grid_row,
            ).font = Font(bold=True)
            for grid_column in range(grid_size):
                value = float(block.values[grid_row, grid_column])
                for data_column in (2, global_data_column):
                    cell = worksheet.cell(
                        first_data_row + grid_row,
                        data_column + grid_column,
                        value,
                    )
                    cell.number_format = "0.000000E+00"

        last_data_row = first_data_row + grid_size - 1
        local_range = (
            f"B{first_data_row}:{get_column_letter(grid_size + 1)}{last_data_row}"
        )
        global_range = (
            f"{get_column_letter(global_data_column)}{first_data_row}:"
            f"{get_column_letter(global_data_column + grid_size - 1)}{last_data_row}"
        )
        local_limits = _color_scale_limits(block.values)
        _add_color_scale(worksheet, local_range, local_limits)
        _add_color_scale(worksheet, global_range, global_limits)

        note_row = last_data_row + 1
        worksheet.cell(note_row, 1, "Farbskala Bild").font = Font(bold=True)
        worksheet.cell(
            note_row,
            2,
            f"{local_limits[0]:.6E} bis {local_limits[2]:.6E}",
        )
        worksheet.cell(
            note_row,
            global_label_column,
            "Farbskala Tabelle",
        ).font = Font(bold=True)
        worksheet.cell(
            note_row,
            global_data_column,
            f"{global_limits[0]:.6E} bis {global_limits[2]:.6E}",
        )
        current_row += grid_size + 8


def _write_definitions(workbook):
    if "Definitions" in workbook.sheetnames:
        workbook.remove(workbook["Definitions"])
    worksheet = workbook.create_sheet("Definitions", 2)
    worksheet.append(("Metrik", "Definition"))
    for metric_name, definition in METRIC_DEFINITIONS:
        worksheet.append((metric_name, definition))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 110


def save_raw_gradient_values(
    accumulator: RawGradientAccumulator,
    sampler_name: str,
    group_name: str,
    excel_path: Path | None = None,
) -> Path | None:
    """Update the group's RawGradientValues workbook for one sampler."""
    if not accumulator.has_data:
        return None

    if excel_path is None:
        output_dir = RESULTS_DIR / "Stats" / "RawGradientValues"
        output_dir.mkdir(parents=True, exist_ok=True)
        excel_path = output_dir / f"{group_name}.xlsx"
    else:
        excel_path = Path(excel_path)
        excel_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(excel_path) if excel_path.exists() else Workbook()
    if not excel_path.exists():
        workbook.remove(workbook.active)

    summary_rows = [
        row
        for row in _read_summary_rows(workbook)
        if str(row.get("Samplername")) != sampler_name
    ]
    summary_rows.extend(_summary_rows_for_sampler(sampler_name, accumulator))
    _write_summary(workbook, summary_rows)

    grid_blocks = _read_grid_blocks(workbook)
    grid_blocks.pop(sampler_name, None)
    if (
        accumulator.first_sample_grid is not None
        and accumulator.first_sample_gradient_shape is not None
        and accumulator.first_index is not None
    ):
        grid_blocks[sampler_name] = FirstSampleGrid(
            sampler_name=sampler_name,
            dataset_index=accumulator.first_index,
            gradient_shape=accumulator.first_sample_gradient_shape,
            values=accumulator.first_sample_grid,
        )
    _write_grid_blocks(workbook, grid_blocks)
    _write_definitions(workbook)

    workbook.save(excel_path)
    print(f"  -> Raw gradient values saved to {excel_path}.")
    return excel_path
